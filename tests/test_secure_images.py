"""Image sources of the authenticated export path.

Public URLs are served by a local TLS server with its own test CA, so the real
connection code runs (address pinning, certificate check for the name,
redirects, limits) inside the network-less test container. Only the resolver,
the CA bundle, the port and the address policy for 127.0.0.1 are replaced.
"""
import asyncio
import base64
from contextlib import ExitStack
import datetime
import http.server
import io
import ipaddress
import os
import socket
import ssl
import struct
import tempfile
import threading
import time
import unittest
from unittest.mock import Mock, patch
import zipfile
import zlib

from test_secure_exports import HEADERS, tool
from tools import secure_images as images
from tools.file_security import FilePolicyError
from cryptography import x509
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import ec
from cryptography.x509.oid import ExtendedKeyUsageOID, NameOID
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

TOKEN = HEADERS["authorization"]


def picture(fmt="PNG", size=(64, 48), mode="RGB", color=(200, 30, 30)):
    buf = io.BytesIO()
    Image.new(mode, size, color).save(buf, fmt)
    return buf.getvalue()


def png_header(width, height):
    """A PNG whose header claims a size; the pixel data is never decoded."""
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    ihdr = struct.pack(">IIBBBBB", width, height, 1, 0, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(b"\0" * 16)) + chunk(b"IEND", b"")


PNG = picture()
NAMES = ["images.test", "images.unsplash.com", "images.pexels.com", "other.test"]


class ImageServer:
    """HTTPS on 127.0.0.1 with a certificate for NAMES, signed by a throwaway CA."""

    def __init__(self):
        self.dir = tempfile.mkdtemp()
        now = datetime.datetime.now(datetime.timezone.utc)
        ca_key = ec.generate_private_key(ec.SECP256R1())
        ca_name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "HC image test CA")])
        ca = (x509.CertificateBuilder().subject_name(ca_name).issuer_name(ca_name)
              .public_key(ca_key.public_key()).serial_number(x509.random_serial_number())
              .not_valid_before(now - datetime.timedelta(hours=1)).not_valid_after(now + datetime.timedelta(days=1))
              .add_extension(x509.BasicConstraints(ca=True, path_length=None), critical=True)
              .add_extension(x509.KeyUsage(False, False, False, False, False, True, True, False, False), critical=True)
              .add_extension(x509.SubjectKeyIdentifier.from_public_key(ca_key.public_key()), critical=False)
              .sign(ca_key, hashes.SHA256()))
        key = ec.generate_private_key(ec.SECP256R1())
        cert = (x509.CertificateBuilder().subject_name(x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, NAMES[0])]))
                .issuer_name(ca_name).public_key(key.public_key()).serial_number(x509.random_serial_number())
                .not_valid_before(now - datetime.timedelta(hours=1)).not_valid_after(now + datetime.timedelta(days=1))
                .add_extension(x509.SubjectAlternativeName([x509.DNSName(n) for n in NAMES]), critical=False)
                .add_extension(x509.BasicConstraints(ca=False, path_length=None), critical=True)
                .add_extension(x509.ExtendedKeyUsage([ExtendedKeyUsageOID.SERVER_AUTH]), critical=False)
                .add_extension(x509.AuthorityKeyIdentifier.from_issuer_public_key(ca_key.public_key()), critical=False)
                .sign(ca_key, hashes.SHA256()))
        self.ca_path = os.path.join(self.dir, "ca.pem")
        cert_path, key_path = os.path.join(self.dir, "cert.pem"), os.path.join(self.dir, "key.pem")
        with open(self.ca_path, "wb") as f:
            f.write(ca.public_bytes(serialization.Encoding.PEM))
        with open(cert_path, "wb") as f:
            f.write(cert.public_bytes(serialization.Encoding.PEM))
        with open(key_path, "wb") as f:
            f.write(key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                                      serialization.NoEncryption()))
        context = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
        context.load_cert_chain(cert_path, key_path)
        self.sni, self.requests = [], []
        context.sni_callback = lambda sock, name, ctx: self.sni.append(name)
        self.httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _handler(self))
        self.httpd.daemon_threads = True
        self.httpd.socket = context.wrap_socket(self.httpd.socket, server_side=True)
        self.port = self.httpd.server_address[1]
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()

    def close(self):
        self.httpd.shutdown()
        self.httpd.server_close()


def _handler(server):
    class Handler(http.server.BaseHTTPRequestHandler):
        protocol_version = "HTTP/1.1"

        def log_message(self, *args):
            pass

        def send(self, status, body=b"", content_type="image/png", **headers):
            self.send_response(status)
            if content_type:
                self.send_header("Content-Type", content_type)
            for name, value in headers.items():
                self.send_header(name.replace("_", "-"), value)
            if "Content_Length" not in headers:
                self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            try:
                self.wfile.write(body)
            except OSError:
                pass

        def redirect(self, location):
            self.send(302, content_type=None, Location=location)

        def do_GET(self):
            server.requests.append((self.headers.get("Host"), self.path))
            path = self.path.split("?")[0]
            formats = {"/ok.png": ("PNG", "image/png"), "/photo.jpg": ("JPEG", "image/jpeg"),
                       "/anim.gif": ("GIF", "image/gif"), "/pic.webp": ("WEBP", "image/webp"),
                       "/pic.bmp": ("BMP", "image/bmp"), "/octet": ("PNG", "application/octet-stream")}
            if path in formats:
                fmt, content_type = formats[path]
                return self.send(200, picture(fmt), content_type)
            if path == "/alpha.png":
                return self.send(200, picture("PNG", mode="RGBA", color=(0, 0, 255, 128)))
            if path == "/huge.png":
                return self.send(200, picture("PNG", size=(4000, 1000)))
            if path == "/page":
                return self.send(200, b"<html>not an image</html>", "text/html")
            if path == "/vector.svg":
                return self.send(200, b"<svg xmlns='http://www.w3.org/2000/svg'/>", "image/svg+xml")
            if path == "/fake.png":
                return self.send(200, b"<html>pretending</html>")
            if path == "/bomb.png":
                return self.send(200, png_header(8000, 6000))
            if path == "/big":
                self.send_response(200)
                self.send_header("Content-Type", "image/png")
                self.send_header("Content-Length", str(images.MAX_BYTES + 1))
                self.end_headers()
                return
            if path in ("/chunked.png", "/big-chunked"):
                self.send_response(200)
                self.send_header("Content-Type", "image/png")
                self.send_header("Transfer-Encoding", "chunked")
                self.end_headers()
                if path == "/chunked.png":
                    body = picture("PNG")
                    pieces = [body[i:i + 100] for i in range(0, len(body), 100)]
                else:
                    pieces = [b"\0" * 65536] * (images.MAX_BYTES // 65536 + 4)
                try:
                    for piece in pieces:
                        self.wfile.write(b"%x\r\n%s\r\n" % (len(piece), piece))
                    self.wfile.write(b"0\r\n\r\n")
                except OSError:
                    pass
                return
            if path in ("/slow-body", "/slow-headers"):
                try:
                    if path == "/slow-body":
                        self.wfile.write(b"HTTP/1.1 200 OK\r\nContent-Type: image/png\r\nContent-Length: 200\r\n\r\n")
                    else:
                        self.wfile.write(b"HTTP/1.1 200 OK\r\n")
                    for _ in range(200):
                        self.wfile.write(b"X" if path == "/slow-headers" else b"\0")
                        self.wfile.flush()
                        time.sleep(0.05)
                except OSError:
                    pass
                self.close_connection = True
                return
            if path == "/gzip":
                return self.send(200, zlib.compress(PNG), Content_Encoding="gzip")
            if path == "/redirect-ok":
                return self.redirect("/ok.png")
            if path == "/redirect-private":
                return self.redirect("https://private.test/ok.png")
            if path == "/redirect-http":
                return self.redirect("http://images.test/ok.png")
            if path == "/redirect-port":
                return self.redirect("https://images.test:8443/ok.png")
            if path == "/redirect-loop":
                return self.redirect("/redirect-loop")
            if path == "/redirect-other-host":
                return self.redirect("https://other.test/ok.png")
            return self.send(404, b"missing", "text/plain")
    return Handler


class Response:
    """Enough of requests.Response for the Open WebUI and search fakes."""

    def __init__(self, status, payload=None, body=b"", headers=None):
        self.status_code, self._payload, self.body = status, payload, body
        self.headers = headers or {}

    def json(self):
        if self._payload is None:
            raise ValueError("no json")
        return self._payload

    def iter_content(self, size):
        for start in range(0, len(self.body), size):
            yield self.body[start:start + size]

    def close(self):
        pass


class OpenWebUI:
    """Files of two users behind the session user and file endpoints."""

    FILES = {
        "own": ("user-1", "image/png", PNG),
        "own-jpeg": ("user-1", "image/jpeg", picture("JPEG")),
        "foreign": ("user-2", "image/png", PNG),
        "report": ("user-1", "application/pdf", b"%PDF-1.4"),
        "broken": ("user-1", "image/png", b"not an image at all"),
        "big": ("user-1", "image/png", b"\0" * 64),
    }

    def __init__(self, session=200):
        self.calls, self.session = [], session

    def __call__(self, url, headers=None, stream=False, **kwargs):
        self.calls.append(url)
        assert headers["Authorization"] == TOKEN, "the forwarded user token must be used"
        if url == "http://webui:8080/api/v1/auths/":
            return Response(self.session, {"id": "user-1"} if self.session == 200 else {"detail": "no"})
        fid, _, rest = url.removeprefix("http://webui:8080/api/v1/files/").partition("/")
        if fid not in self.FILES:
            return Response(404, {"detail": "Not found"})
        owner, content_type, body = self.FILES[fid]
        if rest == "content":
            assert stream, "file content must be streamed"
            return Response(200, body=body, headers={"Content-Length": str(len(body))})
        return Response(200, {"id": fid, "user_id": owner, "meta": {"content_type": content_type}})


class Rules(unittest.TestCase):
    """Rules that need no network at all."""

    def test_address_policy(self):
        refused = ["10.0.0.1", "172.16.0.1", "172.31.255.255", "192.168.0.1", "127.0.0.1", "169.254.169.254",
                   "100.64.0.1", "0.0.0.0", "224.0.0.1", "255.255.255.255", "198.18.0.1", "192.0.2.1",
                   "::1", "::", "fe80::1", "fc00::1", "fd12:3456::1", "ff02::1", "fec0::1", "::ffff:10.0.0.1",
                   "64:ff9b::a00:1", "2002:a00:1::1", "2001:0:4136:e378:8000:63bf:3fff:fdd2", "2001:db8::1"]
        public = ["1.1.1.1", "8.8.8.8", "93.184.216.34", "2606:4700:4700::1111", "64:ff9b::808:808", "::ffff:8.8.8.8"]
        for value in refused:
            with self.subTest(value=value):
                self.assertFalse(images._is_public_ip(ipaddress.ip_address(value)))
        for value in public:
            with self.subTest(value=value):
                self.assertTrue(images._is_public_ip(ipaddress.ip_address(value)))

    def test_loopback_names_are_refused_by_the_real_resolver(self):
        for host in ("localhost", "127.0.0.1", "::1"):
            with self.subTest(host=host), self.assertRaises(images._Unavailable):
                images._public_addresses(host)

    def test_links_that_are_refused_before_any_lookup(self):
        cases = {
            "/etc/passwd": "Local file paths", "../secret.png": "Local file paths", "logo.png": "Local file paths",
            r"C:\Users\x\logo.png": "Local file paths", "file:///etc/passwd": "Local file paths",
            "http://example.com/a.png": "not supported", "ftp://example.com/a.png": "not supported",
            "javascript:alert(1)": "not supported", "gopher://example.com/": "not supported",
            "//example.com/a.png": "must start with https://", "https://exa mple.com/a.png": "spaces",
            "https://example.com/" + "a" * 2100: "2048", "image_query:  ": "needs a query", "": "needs a source",
        }
        with patch.object(images, "_resolve") as resolve, patch.object(images, "_PinnedConnection") as connect:
            for link, message in cases.items():
                with self.subTest(link=link[:40]), self.assertRaises(FilePolicyError) as caught:
                    images.classify(link)
                self.assertIn(message, str(caught.exception))
            with images.request_scope(TOKEN):
                for link, message in {"https://user:pw@images.test/ok.png": "user names",
                                      "https://images.test:8443/ok.png": "standard https port"}.items():
                    with self.subTest(link=link), self.assertRaises(FilePolicyError) as caught:
                        images.load(*images.classify(link))
                    self.assertIn(message, str(caught.exception))
            resolve.assert_not_called()
            connect.assert_not_called()

    def test_sources_of_blocks_and_slides(self):
        self.assertEqual(images.source_of({"file_id": " abc-1 "}), ("file", "abc-1"))
        self.assertEqual(images.source_of({"url": "/api/v1/files/abc/content"}), ("file", "abc"))
        self.assertEqual(images.source_of({"query": "harbour"}), ("query", "harbour"))
        self.assertEqual(images.classify(images.link("query", "a b/c?")), ("query", "a b/c?"))
        for item in ({}, {"file_id": "a", "url": "https://x.test/a.png"}, {"query": ""}, {"url": 5},
                     {"file_id": "../a"}, {"url": "image_query:cats"}):
            with self.subTest(item=item), self.assertRaises(FilePolicyError):
                images.source_of(item)
        with self.assertRaises(FilePolicyError):
            images.slide_source({"image_url": "https://x.test/a.png", "image_query": "cats"})
        self.assertIsNone(images.slide_source({"title": "no image"}))

    def test_data_urls(self):
        good = "data:image/png;base64," + base64.b64encode(PNG).decode()
        with images.request_scope(TOKEN):
            self.assertEqual(images.load("data", good).ext, "png")
            for link, message in {"data:image/svg+xml;base64,PHN2Zy8+": "base64 data: URL",
                                  "data:text/html;base64,PGh0bWw+": "base64 data: URL",
                                  "data:image/png;base64,***": "base64 data: URL",
                                  "data:image/png;base64,QUJD=": "not valid base64",
                                  "data:image/png;base64," + base64.b64encode(b"plain text").decode(): "not a supported image"}.items():
                with self.subTest(link=link[:30]), self.assertRaises(FilePolicyError) as caught:
                    images.load("data", link)
                self.assertIn(message, str(caught.exception))
            with patch.object(images, "MAX_BYTES", 10), self.assertRaises(FilePolicyError) as caught:
                images.load("data", good)
            self.assertIn("larger than", str(caught.exception))

    def test_decoding_limits(self):
        with images.request_scope(TOKEN):
            with self.assertRaises(FilePolicyError) as caught:
                images.load("data", "data:image/png;base64," + base64.b64encode(png_header(8000, 6000)).decode())
            self.assertIn("too many pixels", str(caught.exception))

    def test_budget_per_request(self):
        good = "data:image/png;base64," + base64.b64encode(PNG).decode()
        with patch.object(images, "MAX_COUNT", 2), images.request_scope(TOKEN):
            images.load("data", good)
            images.load("data", good)
            with self.assertRaises(FilePolicyError) as caught:
                images.load("data", good)
            self.assertIn("At most 2 images", str(caught.exception))
        with patch.object(images, "MAX_TOTAL_BYTES", len(PNG) + 10), images.request_scope(TOKEN):
            images.load("data", good)
            with self.assertRaises(FilePolicyError) as caught:
                images.load("data", good)
            self.assertIn("total at most", str(caught.exception))


class PublicImages(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ImageServer()

    @classmethod
    def tearDownClass(cls):
        cls.server.close()

    def setUp(self):
        self.addresses = {name: ["127.0.0.1"] for name in NAMES}
        self.addresses.update({"private.test": ["10.0.0.7"], "mixed.test": ["127.0.0.1", "192.168.1.9"],
                               "v6private.test": ["fd00::1"], "wrongname.test": ["127.0.0.1"]})
        self.resolved = []
        real = images._is_public_ip

        def resolve(host, port):
            self.resolved.append(host)
            if host not in self.addresses:
                raise socket.gaierror("unknown name")
            return self.addresses[host]

        self.stack = ExitStack()
        self.stack.enter_context(patch.object(images, "_resolve", resolve))
        self.stack.enter_context(patch.object(images, "_is_public_ip", lambda ip: str(ip) == "127.0.0.1" or real(ip)))
        self.stack.enter_context(patch.object(images, "_HTTPS_PORT", self.server.port))
        self.stack.enter_context(patch.object(images, "_CA_BUNDLE", self.server.ca_path))
        self.stack.enter_context(images.request_scope(TOKEN))

    def tearDown(self):
        self.stack.close()

    def load(self, link):
        return images.load(*images.classify(link))

    def refused(self, link, message="could not be loaded"):
        with self.assertRaises(FilePolicyError) as caught:
            self.load(link)
        self.assertIn(message, str(caught.exception))
        return str(caught.exception)

    def test_formats_are_decoded_and_re_encoded(self):
        for path, ext in {"/ok.png": "png", "/photo.jpg": "jpeg", "/anim.gif": "png", "/pic.webp": "jpeg",
                          "/pic.bmp": "png", "/octet": "png", "/alpha.png": "png", "/chunked.png": "png"}.items():
            with self.subTest(path=path):
                result = self.load("https://images.test" + path)
                self.assertEqual(result.ext, ext)
                with Image.open(result.stream()) as decoded:
                    self.assertEqual(decoded.format, ext.upper())
                    self.assertEqual(decoded.size, (64, 48))
        with Image.open(self.load("https://images.test/alpha.png").stream()) as decoded:
            self.assertEqual(decoded.mode, "RGBA")
        self.assertEqual(self.load("https://images.test/huge.png").width, images.MAX_SIDE)

    def test_connection_goes_to_the_checked_address_with_tls_for_the_name(self):
        self.resolved.clear()
        self.load("https://images.test/ok.png?size=large")
        self.assertEqual(self.resolved, ["images.test"])
        self.assertEqual(self.server.sni[-1], "images.test")
        self.assertEqual(self.server.requests[-1], (f"images.test:{self.server.port}", "/ok.png?size=large"))
        self.refused("https://wrongname.test/ok.png")

    def test_non_public_addresses_are_refused_before_connecting(self):
        with patch.object(images, "_PinnedConnection") as connect:
            for host in ("private.test", "mixed.test", "v6private.test", "unknown.test"):
                with self.subTest(host=host):
                    message = self.refused(f"https://{host}/ok.png")
                    self.assertNotIn("10.0.0", message)
                    self.assertNotIn("192.168", message)
            connect.assert_not_called()

    def test_redirects_are_checked_again(self):
        self.assertEqual(self.load("https://images.test/redirect-ok").ext, "png")
        with patch.object(images, "_PinnedConnection", wraps=images._PinnedConnection) as connect:
            self.refused("https://images.test/redirect-private")
            self.assertEqual(connect.call_count, 1)
        self.refused("https://images.test/redirect-http", "only https")
        self.refused("https://images.test/redirect-port", "standard https port")
        self.refused("https://images.test/redirect-loop")

    def test_size_time_and_content_limits(self):
        with patch.object(images, "MAX_BYTES", 200_000):
            self.refused("https://images.test/big", "larger than")
            self.refused("https://images.test/big-chunked", "larger than")
        for path in ("/slow-body", "/slow-headers"):
            with self.subTest(path=path), patch.object(images, "TIMEOUT", 1.0):
                started = time.monotonic()
                self.refused("https://images.test" + path)
                self.assertLess(time.monotonic() - started, 4)
        self.refused("https://images.test/gzip")
        self.refused("https://images.test/missing")
        for path in ("/page", "/vector.svg", "/fake.png"):
            with self.subTest(path=path):
                self.refused("https://images.test" + path, "not a supported image")
        self.refused("https://images.test/bomb.png", "too many pixels")

    def test_image_search_stays_on_the_provider_hosts(self):
        found = {"url": "https://images.unsplash.com/ok.png?w=1080"}

        def api(url, params=None, headers=None, **kwargs):
            api.calls.append((url, params, headers))
            if "unsplash" in url:
                return Response(200, {"results": [{"urls": {"regular": found["url"]}}]} if found["url"] else {"results": []})
            return Response(200, {"photos": [{"src": {"large": "https://images.pexels.com/ok.png"}}]})
        api.calls = []

        with patch.object(images, "http_get", api), \
                patch.dict(os.environ, {"IMAGE_SOURCE": "unsplash", "UNSPLASH_ACCESS_KEY": "test-key"}):
            self.assertEqual(images.load("query", "harbour").ext, "png")
            self.assertEqual(api.calls[-1][0], "https://api.unsplash.com/search/photos")
            self.assertEqual(api.calls[-1][1]["query"], "harbour")
            self.assertEqual(api.calls[-1][2], {"Authorization": "Client-ID test-key"})
            for url in ("https://other.test/ok.png", "https://images.unsplash.com/redirect-other-host"):
                found["url"] = url
                with self.subTest(url=url), self.assertRaises(FilePolicyError):
                    images.load("query", "harbour")
            found["url"] = "http://images.unsplash.com/ok.png"
            with self.assertRaises(FilePolicyError) as caught:
                images.load("query", "harbour")
            self.assertIn("only https", str(caught.exception))
            found["url"] = None
            with self.assertRaises(FilePolicyError) as caught:
                images.load("query", "harbour")
            self.assertIn("found no image", str(caught.exception))
        with patch.object(images, "http_get", api), \
                patch.dict(os.environ, {"IMAGE_SOURCE": "pexels", "PEXELS_ACCESS_KEY": "pexels-key"}):
            self.assertEqual(images.load("query", "harbour").ext, "png")
            self.assertEqual(api.calls[-1][2], {"Authorization": "pexels-key"})
        failing = Mock(return_value=Response(500, {}))
        for env, message in (({"IMAGE_SOURCE": "unsplash", "UNSPLASH_ACCESS_KEY": ""}, "not configured"),
                             ({"IMAGE_SOURCE": "local_sd"}, "not available in secure exports"),
                             ({"IMAGE_SOURCE": "unsplash", "UNSPLASH_ACCESS_KEY": "k"}, "unavailable")):
            with self.subTest(env=env), patch.object(images, "http_get", failing), patch.dict(os.environ, env), \
                    self.assertRaises(FilePolicyError) as caught:
                images.load("query", "harbour")
            self.assertIn(message, str(caught.exception))


class OwnFiles(unittest.TestCase):
    def test_only_the_users_own_image_files(self):
        webui = OpenWebUI()
        with patch.object(images, "http_get", webui), images.request_scope(TOKEN):
            self.assertEqual(images.load("file", "own").ext, "png")
            self.assertEqual(images.load("file", "own-jpeg").ext, "jpeg")
            for fid, message in {"foreign": "not found among your files", "missing": "not found among your files",
                                 "report": "is not an image", "broken": "not a supported image"}.items():
                with self.subTest(fid=fid), self.assertRaises(FilePolicyError) as caught:
                    images.load("file", fid)
                self.assertIn(message, str(caught.exception))
            with patch.object(images, "MAX_BYTES", 10), self.assertRaises(FilePolicyError) as caught:
                images.load("file", "big")
            self.assertIn("larger than", str(caught.exception))
            with self.assertRaises(FilePolicyError):
                images.load("file", "../own")
        self.assertEqual(webui.calls.count("http://webui:8080/api/v1/auths/"), 1)
        self.assertFalse(any("foreign/content" in call for call in webui.calls))
        with patch.object(images, "http_get", OpenWebUI(session=401)), images.request_scope(TOKEN), \
                self.assertRaises(FilePolicyError) as caught:
            images.load("file", "own")
        self.assertIn("session could not be confirmed", str(caught.exception))


class ImagesInDocuments(unittest.TestCase):
    def setUp(self):
        self.uploaded = []
        self.stack = ExitStack()
        self.stack.enter_context(patch.object(images, "http_get", OpenWebUI()))
        self.post = self.stack.enter_context(patch.object(tool, "http_post", self.upload))

    def tearDown(self):
        self.stack.close()

    def upload(self, url, **kwargs):
        name, stream, _ = kwargs["files"]["file"]
        self.uploaded.append((name, stream.read()))
        return Mock(status_code=200, json=lambda: {"id": "owned-file-1"})

    def create(self, data):
        return asyncio.run(tool.create_file(data, mcpo_headers=HEADERS))

    def created(self, data):
        result = self.create(data)
        self.assertTrue(result.get("success"), result)
        return self.uploaded[-1][1]

    def refused(self, data, message):
        before = len(self.uploaded)
        result = self.create(data)
        self.assertFalse(result.get("success"), result)
        self.assertIn(message, result["error"]["message"])
        self.assertEqual(len(self.uploaded), before)
        self.assertEqual(list(os.scandir(tool.EXPORT_DIR)), [])

    def test_pdf(self):
        raw = self.created({"format": "pdf", "content": [
            {"type": "paragraph", "text": "Logo below"}, {"type": "image", "file_id": "own", "alt": "Logo"}]})
        self.assertEqual(sum(len(page.images) for page in PdfReader(io.BytesIO(raw)).pages), 1)
        raw = self.created({"format": "pdf", "content": "Intro ![Logo](/api/v1/files/own/content) outro\n\n"
                                                        "```\n![example](/etc/passwd)\n```"})
        reader = PdfReader(io.BytesIO(raw))
        text = " ".join(page.extract_text() for page in reader.pages)
        self.assertIn("Intro", text)
        self.assertIn("outro", text)
        self.assertIn("![example](/etc/passwd)", text)
        self.assertEqual(sum(len(page.images) for page in reader.pages), 1)
        self.refused({"format": "pdf", "content": "- ![x](/api/v1/files/own/content)"}, "inside lists")
        self.refused({"format": "pdf", "content": "See ![x](https://images.test/a.png and more"}, "could not be read")
        self.refused({"format": "pdf", "content": "![x](/etc/passwd)"}, "Local file paths")
        self.refused({"format": "pdf", "content": [{"type": "image", "file_id": "foreign"}]}, "not found among your files")
        self.refused({"format": "pdf", "content": [{"type": "image", "url": "http://169.254.169.254/latest"}]}, "not supported")

    def test_docx(self):
        raw = self.created({"format": "docx", "content": [{"type": "image", "file_id": "own"},
                                                          "![Logo](/api/v1/files/own-jpeg/content)",
                                                          {"type": "paragraph", "text": "Text"}]})
        self.assertEqual(len(Document(io.BytesIO(raw)).inline_shapes), 2)
        raw = self.created({"format": "docx", "content": "# Title\n![Logo](/api/v1/files/own/content)\nText"})
        document = Document(io.BytesIO(raw))
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertLessEqual(document.inline_shapes[0].width, 6 * 914400)
        self.refused({"format": "docx", "content": [{"type": "paragraph", "text": "Logo ![x](/api/v1/files/own/content)"}]},
                     "its own image block or line")
        self.refused({"format": "docx", "content": "Logo: ![x](/api/v1/files/own/content)"}, "its own image block or line")
        self.refused({"format": "docx", "content": [{"type": "image", "query": "cats", "url": "https://x.test/a.png"}]},
                     "exactly one of file_id, url or query")
        self.refused({"format": "docx", "content": ["Text", {"type": "image", "file_id": "foreign"}]},
                     "not found among your files")

    def test_pptx(self):
        raw = self.created({"format": "pptx", "slides_data": [
            {"title": "Logo", "content": ["Point"], "image_file_id": "own", "image_position": "left", "image_size": "large"}]})
        slide = Presentation(io.BytesIO(raw)).slides[1]
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(pictures), 1)
        self.assertAlmostEqual(pictures[0].width / pictures[0].height, 64 / 48, places=2)
        self.refused({"format": "pptx", "slides_data": [{"title": "x", "content": ["![x](/api/v1/files/own/content)"]}]},
                     "not in the slide text")
        self.refused({"format": "pptx", "slides_data": [{"title": "x", "content": ["y"], "image_file_id": "own",
                                                         "image_query": "cats"}]}, "one of image_file_id")
        self.refused({"format": "pptx", "slides_data": [{"title": "x", "content": ["y"], "image_url": "../x.png"}]},
                     "Local file paths")
        self.refused({"format": "pptx", "slides_data": [{"title": "x", "content": ["y"], "image_file_id": "report"}]},
                     "is not an image")

    def test_archive_with_images_uploads_once(self):
        result = asyncio.run(tool.generate_and_archive([
            {"format": "docx", "filename": "a.docx", "content": [{"type": "image", "file_id": "own"}]},
            {"format": "pptx", "filename": "b.pptx", "slides_data": [{"title": "x", "content": ["y"], "image_file_id": "own"}]},
        ], mcpo_headers=HEADERS))
        self.assertTrue(result.get("success"), result)
        self.assertEqual(len(self.uploaded), 1)
        with zipfile.ZipFile(io.BytesIO(self.uploaded[0][1])) as archive:
            for name, media in (("a.docx", "word/media/"), ("b.pptx", "ppt/media/")):
                with zipfile.ZipFile(io.BytesIO(archive.read(name))) as inner:
                    self.assertTrue(any(n.startswith(media) for n in inner.namelist()), name)

    def test_too_many_images_fail_before_generation(self):
        blocks = [{"type": "image", "file_id": "own"}] * (images.MAX_COUNT + 1)
        self.refused({"format": "docx", "content": blocks}, f"At most {images.MAX_COUNT} images")


if __name__ == "__main__":
    unittest.main()
