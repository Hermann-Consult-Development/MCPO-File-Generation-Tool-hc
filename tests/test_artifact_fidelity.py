"""Real artifact regression checks; pypdf is a test-only dependency."""
import asyncio
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

from test_secure_exports import HEADERS, tool
from tools.spreadsheet_formulas import relocate_formula
from tools.file_security import FilePolicyError
from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader


class ArtifactFidelity(unittest.TestCase):
    def setUp(self):
        self.uploaded = []

    def upload(self, url, **kwargs):
        self.uploaded.append(kwargs['files']['file'][1].read())
        return Mock(status_code=200, json=lambda: {'id': 'owned-artifact'})

    def generate(self, data):
        with patch.object(tool, 'http_post', self.upload):
            return asyncio.run(tool.create_file(data, mcpo_headers=HEADERS))

    def test_grid_and_anchored_formulas_move_with_template(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / 'template.xlsx'
            wb = Workbook()
            wb.active.auto_filter.ref = 'B5:E8'
            wb.save(template)
            data = [['Artikel', 'Menge', 'Preis', 'Gesamt'],
                    ['Äpfel', 3, 12.5, '=B2*C2'],
                    ['Öl', 2, 8.75, "='Kalkulation'!$B$3*C$3"],
                    ['Summe', '=SUM(B2:B3)', '', '=SUM($D$2:$D$3)']]
            with patch.object(tool, 'XLSX_TEMPLATE', wb), patch.object(tool, 'XLSX_TEMPLATE_PATH', str(template)):
                result = self.generate({'format': 'xlsx', 'title': 'Kalkulation', 'content': data})
            self.assertTrue(result['success'], result)
            ws = load_workbook(io.BytesIO(self.uploaded[0])).active
            self.assertEqual(ws['B6'].value, 'Äpfel')
            self.assertEqual(ws['C6'].value, 3)
            self.assertEqual(ws['D6'].value, 12.5)
            self.assertEqual(ws['E6'].value, '=C6*D6')
            self.assertEqual(ws['E7'].value, "='Kalkulation'!$C$7*D$7")
            self.assertEqual(ws['C8'].value, '=SUM(C6:C7)')
            self.assertEqual(ws['E8'].value, '=SUM($E$6:$E$7)')

    def test_formula_token_boundaries_and_existing_other_sheet(self):
        translate = lambda formula: relocate_formula(formula, 4, 1, "Müller's Blatt", ["Müller's Blatt", 'Lookup'])
        self.assertEqual(translate('=IF(A1="A1",$B$2,C$3+$D4)'), '=IF(B5="A1",$C$6,D$7+$E8)')
        self.assertEqual(translate("='Müller''s Blatt'!A1+Lookup!$B$2"), "='Müller''s Blatt'!B5+Lookup!$B$2")
        self.assertEqual(translate('=SUM($B:$C)+SUM($2:$4)'), '=SUM($C:$D)+SUM($6:$8)')
        for formula in ['=NamedTotal', '=Lookup!NamedTotal', '=Table1[Amount]', '=INDIRECT("B2")', '=[External.xlsx]Sheet1!A1', '=Missing!A1', '=Sheet1:Sheet3!A1', '=XFD1048576', '=SUM(A:5)']:
            with self.subTest(formula=formula), self.assertRaises(FilePolicyError):
                translate(formula)

    def test_unsupported_formula_is_not_published(self):
        result = self.generate({'format': 'xlsx', 'content': [['Wert'], ['=UnknownTotal']]})
        self.assertFalse(result['success'])
        self.assertIn('Named ranges', result['error']['message'])
        self.assertEqual(self.uploaded, [])

    def test_tall_pdf_table_cell_spans_pages_without_losing_content(self):
        words = [f'MARKER{i:03d}' for i in range(300)]
        result = self.generate({'format': 'pdf', 'content': [{'type': 'table', 'headers': ['Position', 'Beschreibung'], 'rows': [['1', ' '.join(words)]]}]})
        self.assertTrue(result['success'], result)
        pdf = PdfReader(io.BytesIO(self.uploaded[0]))
        self.assertGreater(len(pdf.pages), 1)
        pages = [p.extract_text() for p in pdf.pages]
        for word in words:
            self.assertIn(word, '\n'.join(pages))
        self.assertTrue(all('Beschreibung' in text for text in pages))
        self.assertNotIn('Error in PDF generation', '\n'.join(pages))

    def test_pdf_layout_error_is_truthful_and_never_uploaded(self):
        # Repeating headers cannot split across pages; preserve a clear failure.
        result = self.generate({'format': 'pdf', 'content': [{'type': 'table', 'headers': ['Überschrift ' * 500], 'rows': [['Daten']]}]})
        self.assertFalse(result['success'])
        self.assertIn('PDF generation failed', result['error']['message'])
        self.assertNotIn('url', result)
        self.assertEqual(self.uploaded, [])

    def test_docx_table_borders_precede_table_look(self):
        result = self.generate({'format': 'docx', 'content': [{'type': 'table', 'data': [['Größe', 'Menge'], ['Äpfel', '3']]}]})
        self.assertTrue(result['success'], result)
        doc = Document(io.BytesIO(self.uploaded[0]))
        tags = [element.tag.rsplit('}', 1)[-1] for element in doc.tables[0]._tbl.tblPr]
        if 'tblBorders' in tags and 'tblLook' in tags:
            self.assertLess(tags.index('tblBorders'), tags.index('tblLook'))
        self.assertEqual(doc.tables[0].cell(1, 0).text, 'Äpfel')


if __name__ == '__main__':
    unittest.main()
