import asyncio
import importlib.util
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch
import zipfile

os.environ.update(REQUIRE_USER_AUTH="true", FILE_EXPORT_DIR=tempfile.mkdtemp(),
                  DOCS_TEMPLATE_DIR="/missing-test-templates", OWUI_URL="http://webui:8080",
                  JWT_SECRET="admin-fallback-must-never-be-used")
from tools import file_export_mcp as tool
from tools import file_security as security
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

HEADERS = {"authorization": "Bearer ordinary-test-user"}
CASES = [
    {"format": "docx", "filename": "Übersicht.docx", "content": [{"type": "paragraph", "text": "HC export"}]},
    {"format": "pdf", "filename": "Übersicht.pdf", "content": [{"type": "paragraph", "text": "HC export"}]},
    {"format": "xlsx", "filename": "Übersicht.xlsx", "content": [["Name", "Value"], ["HC export", "42"]]},
    {"format": "csv", "filename": "Übersicht.csv", "content": [["Name", "Value"], ["HC export", "42"]]},
    {"format": "pptx", "filename": "Übersicht.pptx", "slides_data": [{"title": "HC export", "content": ["Qualified file"]}]},
]


def validate_artifact(kind, raw):
    assert len(raw) > 10
    if kind == "docx":
        assert "HC export" in " ".join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
    elif kind == "xlsx":
        assert load_workbook(io.BytesIO(raw)).active.cell(2, 1).value == "HC export"
    elif kind == "pptx":
        slides = Presentation(io.BytesIO(raw)).slides
        # Existing renderer creates its cover plus one supplied content slide.
        assert len(slides) == 2
        assert any("HC export" in shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame)
    elif kind == "csv":
        assert "HC export,42" in raw.decode("utf-8")
    else:
        assert raw.startswith(b"%PDF-") and raw.rstrip().endswith(b"%%EOF")


class SecureExports(unittest.TestCase):
    def setUp(self):
        self.uploaded = []

    def upload(self, url, **kwargs):
        self.assertEqual(url, "http://webui:8080/api/v1/files/")
        self.assertEqual(kwargs["headers"]["Authorization"], HEADERS["authorization"])
        self.assertEqual(kwargs["params"], {"process": "false"})
        name, stream, _ = kwargs["files"]["file"]
        self.uploaded.append((name, stream.read()))
        return Mock(status_code=200, json=lambda: {"id": "owned-file-1"})

    def test_five_formats_are_real_and_upload_owned(self):
        with patch.object(tool, "http_post", self.upload):
            for case in CASES:
                result = asyncio.run(tool.create_file(case, mcpo_headers=HEADERS))
                self.assertTrue(result.get("success"), result)
                self.assertEqual(result["url"], "/api/v1/files/owned-file-1/content")
                self.assertNotIn("/output", json.dumps(result))
                self.assertNotIn("localhost:9003", json.dumps(result))
                validate_artifact(case["format"], self.uploaded[-1][1])
        self.assertEqual(list(Path(tool.EXPORT_DIR).iterdir()), [])

    def test_archive_contains_all_five_and_uploads_once(self):
        with patch.object(tool, "http_post", self.upload):
            result = asyncio.run(tool.generate_and_archive(CASES, archive_name="Übersicht", mcpo_headers=HEADERS))
        self.assertTrue(result.get("success"), result)
        self.assertEqual(len(self.uploaded), 1)
        with zipfile.ZipFile(io.BytesIO(self.uploaded[0][1])) as archive:
            for case in CASES:
                validate_artifact(case["format"], archive.read(case["filename"]))

    def test_auth_never_uses_admin_fallback(self):
        with patch.object(tool, "http_post") as post:
            result = asyncio.run(tool.create_file(CASES[0]))
            self.assertFalse(result["success"])
            post.assert_not_called()
        self.assertEqual(security.user_token({"Authorization": "Bearer user"}, "admin"), "Bearer user")
        for headers in ({}, {"authorization": ""}, {"authorization": "Bearer "}):
            with self.assertRaises(ValueError):
                security.user_token(headers, "admin")

    def test_paths_rejected_for_all_tool_entries(self):
        for name in ("../outside.docx", "/tmp/out.docx", r"C:\outside.docx", r"..\out.docx", "a/b.docx"):
            with self.subTest(name=name), patch.object(tool, "http_post") as post, patch.object(tool, "http_get") as get:
                outputs = [
                    asyncio.run(tool.create_file(dict(CASES[0], filename=name), mcpo_headers=HEADERS)),
                    asyncio.run(tool.generate_and_archive(CASES, archive_name=name, mcpo_headers=HEADERS)),
                    asyncio.run(tool.generate_and_archive([dict(CASES[0], filename=name)], mcpo_headers=HEADERS)),
                    asyncio.run(tool.full_context_document("file-id", name, mcpo_headers=HEADERS)),
                    asyncio.run(tool.edit_document("file-id", name, [], mcpo_headers=HEADERS)),
                    asyncio.run(tool.review_document("file-id", name, [], mcpo_headers=HEADERS)),
                ]
                for result in outputs: self.assertFalse(result["success"], result)
                post.assert_not_called(); get.assert_not_called()

    def test_upload_failure_is_not_a_download_success(self):
        for code in (301, 401, 500):
            with patch.object(tool, "http_post", return_value=Mock(status_code=code)):
                result = asyncio.run(tool.create_file(CASES[3], mcpo_headers=HEADERS))
                self.assertFalse(result["success"])
                self.assertNotIn("url", result)
        with patch.object(tool, "http_post", side_effect=RuntimeError("/private/path secret-token")):
            result = asyncio.run(tool.create_file(CASES[3], mcpo_headers=HEADERS))
            self.assertNotIn("private", json.dumps(result))
            self.assertNotIn("secret-token", json.dumps(result))

    def test_untrusted_images_fail_explicitly_without_file_or_network_access(self):
        for content in ("![local](/etc/private.png)", '<img src="http://169.254.169.254/latest">',
                        [{"type":"image_query", "query":"anything"}]):
            with patch.object(tool, "http_get") as get, patch.object(tool, "http_post") as post:
                result = asyncio.run(tool.create_file({"format":"pdf", "content":content}, mcpo_headers=HEADERS))
                self.assertFalse(result["success"])
                self.assertIn("not supported", result["error"]["message"])
                get.assert_not_called(); post.assert_not_called()

    def test_http_has_timeouts_and_no_credential_redirect(self):
        for method in ("get", "post"):
            with patch.object(security.requests, method) as request:
                getattr(security, "http_" + method)("http://webui:8080", headers=HEADERS, allow_redirects=True)
                self.assertFalse(request.call_args.kwargs["allow_redirects"])
                self.assertEqual(request.call_args.kwargs["timeout"], (5, 60))

    def test_empty_excel_is_a_real_artifact(self):
        with patch.object(tool, "http_post", self.upload):
            result = asyncio.run(tool.create_file({"format":"xlsx", "content":[]}, mcpo_headers=HEADERS))
        self.assertTrue(result.get("success"), result)
        self.assertTrue(load_workbook(io.BytesIO(self.uploaded[0][1])).sheetnames)

    def test_public_file_server_is_disabled(self):
        from fastapi.testclient import TestClient
        path = Path(__file__).resolve().parents[1] / "LLM_Export/docker/file_server/file_export_server.py"
        spec = importlib.util.spec_from_file_location("export_server_test", path)
        module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
        folder = Path(tool.EXPORT_DIR) / "public"; folder.mkdir(exist_ok=True)
        (folder / "secret.txt").write_text("not public")
        with TestClient(module.app) as client:
            self.assertEqual(client.get("/files/public/secret.txt").status_code, 404)
        (folder / "secret.txt").unlink(); folder.rmdir()


if __name__ == "__main__":
    unittest.main()
